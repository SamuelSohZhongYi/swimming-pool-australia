"""Extract 'Participation by activity (adults)' from AusPlay national data tables.

Usage: python3 -I parse_ausplay.py <raw_dir> <out_csv>
Writes long-format rows: period, series, activity, block, column, value
"""
import glob, csv, re, sys, warnings
import openpyxl

warnings.filterwarnings("ignore")
raw, out = sys.argv[1], sys.argv[2]
TARGET = "participation by activity (adults)"


def table_sheet(wb):
    """Find the sheet number for the target table via the Index sheet."""
    ws = wb["Index"]
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row if v is not None]
        for i, v in enumerate(vals):
            if isinstance(v, str) and v.strip().lower() == TARGET and i > 0:
                num = str(vals[i - 1]).strip()
                if num in wb.sheetnames:
                    return num  # first match = the main (not year-on-year) table
    return None


def period_of(wb):
    ws = wb["Index"]
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and v.startswith("AusPlay survey results"):
                return v.replace("AusPlay survey results", "").strip()
    return "?"


rows_out = []
for f in sorted(glob.glob(f"{raw}/*National*.xlsx")):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    sheet, period = table_sheet(wb), period_of(wb)
    ws = wb[sheet]
    data = [list(r) for r in ws.iter_rows(values_only=True)]
    # column header row = the row containing the 15-17 age band
    h = next(i for i, r in enumerate(data) if any(isinstance(x, str) and x.strip() == "15-17" for x in r))
    top, sub = data[h - 1], data[h]
    cols, current_top = {}, None
    for j in range(1, len(sub)):
        if top[j] not in (None, " "):
            current_top = str(top[j]).strip()
        s = sub[j]
        if s in (None, " "):
            continue
        s = str(s).strip()
        if s.startswith("Total"):
            name = "Total" if (current_top in (None, "Total", "Total#", "Age group") or current_top.startswith("Total")) else current_top
        else:
            name = s
        if name in ("Males", "Females") or s in ("Males", "Females"):
            name = s if s in ("Males", "Females") else name
        cols[j] = name.replace("#", "")
    block = None
    for r in data[h + 1:]:
        if r[0] is None and isinstance(r[1], str) and r[1].strip():
            block = r[1].strip()
            continue
        if isinstance(r[0], str) and r[0].strip() and block:
            name = r[0].strip()
            if name.lower().startswith(("note", "source", "#", "*", "^")):
                continue
            for j, c in cols.items():
                v = r[j] if j < len(r) else None
                if isinstance(v, (int, float)):
                    rows_out.append([period, "online" if int(re.findall(r"20\d\d", period)[-1]) >= 2024 else "phone",
                                     name, block, c, v])

with open(out, "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["period", "series", "activity", "block", "column", "value"])
    w.writerows(rows_out)
print(len(rows_out), "rows written")
